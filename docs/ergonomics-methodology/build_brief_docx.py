#!/usr/bin/env python3
"""Собирает DOCX из 2026-10-19_consultation-brief.md.

Формат: A4, поля 1,8/1,8/2,0/2,0 см, Times New Roman 11 pt,
интервал 1,15, абзацы без отступа, таблицы 8–9 pt.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
MD = HERE / "2026-10-19_consultation-brief.md"
OUT = HERE / "2026-10-19_consultation-brief.docx"
FONT = "Times New Roman"


def set_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def add_inline(paragraph, text: str, size: float, bold: bool = False) -> None:
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
            run.bold = bold
        set_font(run, size)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    m = re.match(r"\*\*(Таблица\s+\d+\.)\*\*\s*(.*)", text)
    if m:
        r = p.add_run(m.group(1) + " ")
        r.bold = True
        set_font(r, 10)
        r = p.add_run(m.group(2))
        set_font(r, 10)
    else:
        add_inline(p, text, 10)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = True
    set_font(run, 12.5)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    add_inline(p, text, 11)


def set_cell_margins(cell, top=14, bottom=14, left=40, right=40) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = tc_mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def shade(cell, fill: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table_layout(table, widths_cm: list[float]) -> None:
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) / 2.54 * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)


def add_table(doc: Document, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        if re.match(r"^\|[\s\-:|]+\|$", line):
            continue
        rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    ncols = max(len(r) for r in rows)
    for row in rows:
        row.extend([""] * (ncols - len(row)))
    header = rows[0]

    if ncols == 9:
        widths = [2.6, 3.0, 1.5, 1.3, 1.5, 1.4, 1.4, 1.5, 2.8]
        centers = {2, 3, 4, 5, 6, 7, 8}
        size = 8.2
    elif ncols == 12:
        widths = [3.0] + [1.27] * 11
        centers = set(range(1, 12))
        size = 8.0
    elif ncols == 5:
        widths = [6.0, 2.5, 2.5, 2.5, 3.5]
        centers = {1, 2, 3, 4}
        size = 8.8
    elif ncols == 4 and header[0].lower().startswith("характеристика"):
        widths = [3.0, 6.0, 2.5, 5.5]
        centers = {2}
        size = 8.5
    elif ncols == 4:
        widths = [7.0, 3.3, 3.3, 3.4]
        centers = {1, 2, 3}
        size = 9.0
    else:
        widths = [17.0 / ncols] * ncols
        centers = set(range(1, ncols))
        size = 8.5

    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_layout(table, widths)

    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if r == 0 or c in centers
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(p, text, size, bold=(r == 0))
        tr_pr = table.rows[r]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if r == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
            for c in range(ncols):
                shade(table.cell(0, c))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("— ")
    set_font(run, 9)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    set_font(run, 9)
    run = paragraph.add_run(" —")
    set_font(run, 9)


def main() -> None:
    lines = MD.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise SystemExit("Нет заголовка # ...")
    title = lines[0][2:].strip()

    body_start = next(i for i, line in enumerate(lines) if line.startswith("## "))
    preamble = [line.strip() for line in lines[1:body_start] if line.strip()]
    subtitle = ""
    authors: list[str] = []
    for line in preamble:
        if line.startswith("**") and line.endswith("**") and not authors:
            subtitle = line[2:-2]
        else:
            authors.append(line)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    rpr.append(lang)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(title)
    run.bold = True
    set_font(run, 15)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(subtitle)
        run.bold = True
        set_font(run, 11)

    if authors:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        for i, line in enumerate(authors):
            if i:
                p.add_run().add_break()
            run = p.add_run(line)
            set_font(run, 10.5)

    i = body_start
    pending_caption: str | None = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip())
            i += 1
            continue
        if line.startswith("**Таблица"):
            pending_caption = line
            add_caption(doc, line)
            i += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            pending_caption = None
            continue
        paragraph = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("## ", "|", "**Таблица")):
                break
            paragraph.append(nxt)
            i += 1
        add_body(doc, " ".join(paragraph))

    add_page_number(section.footer.paragraphs[0])

    doc.core_properties.title = title
    doc.core_properties.author = "Папин А.В."
    doc.core_properties.subject = "Выбор модели машинного обучения"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
