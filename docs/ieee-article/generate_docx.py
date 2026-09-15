#!/usr/bin/env python3
"""Генерация .docx статьи PIERE 2026 из Markdown по шаблону IEEE.

Читает `draft-article-en.md`, создаёт `PIERE2026_Papin_generated.docx`:
заголовок/авторы/аннотация — в одну колонку, тело — в две колонки,
Times New Roman, таблицы, рисунки, подписи. Абзацы собираются из
непрерывных строк (по пустой строке).

Зависимости: python-docx (pip install python-docx).

Запуск:
    python3 docs/ieee-article/generate_docx.py
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "draft-article-en.md")
OUT = os.path.join(HERE, "PIERE2026_Papin_generated.docx")

ALIGN = {
    "just": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
}


def set_cols(section, num):
    cols = section._sectPr.xpath("./w:cols")
    if cols:
        cols[0].set(qn("w:num"), str(num))


def add(doc, text, size=10, bold=False, italic=False, align="just", sb=0, sa=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    p.alignment = ALIGN[align]
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    return p


def flush(doc, buf, first_body_ref):
    """Записать накопленный абзац."""
    if not buf:
        return
    text = " ".join(buf)
    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    if text.strip():
        add(doc, text.strip())
    buf.clear()


def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.78)
    sec.top_margin = sec.bottom_margin = Cm(1.9)

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    buf = []
    body_started = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Таблица — сначала сбросить абзац
        if line.startswith("|"):
            flush(doc, buf, None)
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[-: ]+$", "".join(cells)):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, c in enumerate(row[: len(rows[0])]):
                        pr = t.cell(ri, ci).paragraphs[0]
                        rr = pr.add_run(c)
                        rr.font.size = Pt(8)
                        rr.font.name = "Times New Roman"
                        rr.bold = ri == 0
            continue

        # Рисунок
        m = re.match(r"!\[.*?\]\((.+?)\)", line)
        if m:
            flush(doc, buf, None)
            path = m.group(1)
            if not os.path.isabs(path):
                path = os.path.join(HERE, path)
            try:
                doc.add_picture(path, width=Cm(7.5))
            except Exception as e:
                add(doc, f"[image: {m.group(1)}] ({e})")
            i += 1
            continue

        if line.startswith("# "):
            flush(doc, buf, None)
            add(doc, line[2:], size=24, align="center", sa=6)
        elif line.startswith("## "):
            flush(doc, buf, None)
            if not body_started:
                new = doc.add_section(WD_SECTION.CONTINUOUS)
                set_cols(new, 2)
                new.left_margin = new.right_margin = Cm(1.78)
                body_started = True
            add(doc, line[3:], size=10, bold=True, align="center", sb=6, sa=2)
        elif line.startswith("### "):
            flush(doc, buf, None)
            add(doc, line[4:], size=10, italic=True, align="left", sb=4, sa=1)
        elif line.startswith("**Fig."):
            flush(doc, buf, None)
            add(doc, re.sub(r"\*\*", "", line), size=8, align="center")
        elif line.startswith("---"):
            flush(doc, buf, None)
        elif not line:
            flush(doc, buf, None)
        else:
            buf.append(line)
        i += 1
    flush(doc, buf, None)

    doc.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
